"""Excel workbook generated FROM the DuckDB package, so the two can never disagree.

One sheet per table, guide sheets first. Excel limits are handled, not hidden: long cells are truncated and
oversized tables cut off, and every such case is listed on the ``xlsx_notes`` sheet.
"""
from __future__ import annotations

import re
from pathlib import Path

from openpyxl import Workbook

from .schema import PACKAGE_TABLES, TABLES

MAX_ROWS = 1_048_575  # Excel sheet limit minus the header row
CELL_LIMIT = 32_000   # Excel cell limit is 32,767 characters
_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_FIRST = ("manifest", "guide", "relationships", "data_dictionary")
_DERIVED = ("encounter_metadata", "image_metadata_values")


def _select(columns: list[tuple[str, str]], cell_limit: int) -> str:
    parts = []
    for name, type_ in columns:
        quoted = f'"{name}"'
        if type_ == "TIMESTAMP WITH TIME ZONE":
            parts.append(f"timezone('UTC', {quoted}) AS {quoted}")  # Excel has no time zones: UTC, naive
        elif type_.endswith("]") or type_.startswith(("STRUCT", "MAP", "UNION")):
            parts.append(f"left(to_json({quoted})::VARCHAR, {cell_limit + 1}) AS {quoted}")  # lists/objects -> JSON text
        elif type_ in {"VARCHAR", "JSON"}:
            parts.append(f"left({quoted}::VARCHAR, {cell_limit + 1}) AS {quoted}")  # +1 so truncation is detectable
        else:
            parts.append(quoted)
    return ", ".join(parts)


def _clean(value, cell_limit: int):
    if isinstance(value, str):
        value = _ILLEGAL.sub("", value)
        if len(value) > cell_limit:
            return value[:cell_limit], True
    return value, False


def write_workbook(duckdb_path: str | Path, xlsx_path: str | Path, *, max_rows: int = MAX_ROWS,
                   cell_limit: int = CELL_LIMIT) -> list[tuple[str, str]]:
    """Write ``xlsx_path`` from the package at ``duckdb_path``; returns the (table, note) pairs also written to the file."""
    import duckdb

    con = duckdb.connect(str(duckdb_path), read_only=True)
    present = {row[0] for row in con.execute("select table_name from information_schema.tables").fetchall()}
    order = list(_FIRST) + [t.name for t in TABLES] + [n for n in _DERIVED]
    order = [name for name in dict.fromkeys(order) if name in present]
    workbook = Workbook(write_only=True)
    notes: list[tuple[str, str]] = []
    for name in order:
        columns = [(row[0], row[1]) for row in con.execute(f'DESCRIBE "{name}"').fetchall()]
        total = con.execute(f'SELECT count(*) FROM "{name}"').fetchone()[0]
        sheet = workbook.create_sheet(title=name[:31])
        sheet.append([column for column, _ in columns])
        cursor = con.execute(f'SELECT {_select(columns, cell_limit)} FROM "{name}" LIMIT {max_rows}')
        truncated = 0
        while chunk := cursor.fetchmany(5000):
            for row in chunk:
                cleaned = []
                for value in row:
                    value, cut = _clean(value, cell_limit)
                    truncated += cut
                    cleaned.append(value)
                sheet.append(cleaned)
        if total > max_rows:
            notes.append((name, f"only the first {max_rows} of {total} rows are in this sheet; the DuckDB file has all rows"))
        if truncated:
            notes.append((name, f"{truncated} cell(s) longer than {cell_limit} characters were truncated"))
    con.close()
    notes_sheet = workbook.create_sheet(title="xlsx_notes")
    notes_sheet.append(["table", "note"])
    for note in notes or [("", "No Excel limits were hit; every row and cell is complete.")]:
        notes_sheet.append(list(note))
    workbook.save(str(xlsx_path))
    return notes
