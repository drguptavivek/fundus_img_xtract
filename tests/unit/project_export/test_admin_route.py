import time
from datetime import date

import duckdb
import pytest

from models import EncounterSetImage, PatientEncounters, Project, User
from project_configuration.models import ProjectLabUnit
from tests.conftest import create_authenticated_client
from tests.helpers.factories import UserFactory


@pytest.fixture
def world(db_session, core_test_data):
    lab = core_test_data["lab_a1"]
    project = Project(title="Route Project", code="ROUTE_P", active=True)
    db_session.add(project)
    db_session.flush()
    db_session.add(ProjectLabUnit(project_id=project.id, lab_unit_id=lab.id, active=True))
    encounter = PatientEncounters(
        name="Jane Patient", patient_id="PID-9", capture_date="2026-09-14", capture_date_dt=date(2026, 9, 14),
        lab_unit_id=lab.id, project_id=project.id, disease_id=core_test_data["dr"].id, is_set_based=True,
    )
    db_session.add(encounter)
    db_session.flush()
    db_session.add(EncounterSetImage(
        patient_encounter_id=encounter.id, spatial_position=1, original_filename="a.jpg",
        folder_rel="files/test", project_id=project.id,
    ))
    admin = UserFactory.create_admin(db_session, username="export_admin")
    plain = User(username="export_plain", password_hash="x", is_active=True, email="plain@example.org")
    db_session.add(plain)
    db_session.flush()
    return {"project": project, "admin": admin, "plain": plain}


def _client(app, db_session, user, *, sudo=True):
    from flask import g

    g.pop("_login_user", None)
    client = create_authenticated_client(app, user, db_session)
    if sudo:
        with client.session_transaction() as sess:
            sess["last_sudo_time"] = int(time.time())
    return client


def _open(tmp_path, data, name="out.duckdb"):
    path = tmp_path / name
    path.write_bytes(data)
    return duckdb.connect(str(path), read_only=True)


def test_admin_downloads_package_without_identifiers_by_default(app, db_session, world, tmp_path):
    client = _client(app, db_session, world["admin"])
    response = client.get(f"/api/admin/projects/{world['project'].id}/export.duckdb")
    assert response.status_code == 200
    assert response.headers["Cache-Control"] == "no-store"
    assert "attachment" in response.headers["Content-Disposition"] and ".duckdb" in response.headers["Content-Disposition"]
    con = _open(tmp_path, response.data)
    assert con.execute("select patient_id, identifiers_included from encounters").fetchone() == (None, False)
    assert con.execute("select count(*) from images").fetchone() == (1,)


def test_identifiers_only_when_asked(app, db_session, world, tmp_path):
    client = _client(app, db_session, world["admin"])
    response = client.get(f"/api/admin/projects/{world['project'].id}/export.duckdb?identifiers=1")
    assert response.status_code == 200
    assert "_identifiers_" in response.headers["Content-Disposition"]
    con = _open(tmp_path, response.data)
    assert con.execute("select patient_id, patient_name from encounters").fetchone() == ("PID-9", "Jane Patient")


def test_non_admin_is_forbidden_and_stale_password_redirects(app, db_session, world):
    url = f"/api/admin/projects/{world['project'].id}/export.duckdb"
    assert _client(app, db_session, world["plain"]).get(url).status_code == 403
    stale = _client(app, db_session, world["admin"], sudo=False).get(url)
    assert stale.status_code == 302 and "confirm" in stale.headers["Location"]


def test_unknown_project_is_404(app, db_session, world):
    response = _client(app, db_session, world["admin"]).get("/api/admin/projects/999999/export.duckdb")
    assert response.status_code == 404 and response.get_json()["error"] == "not_found"


def test_no_project_export_multi_select_and_validation(app, db_session, world, core_test_data, tmp_path):
    lab_a1, lab_a2 = core_test_data["lab_a1"], core_test_data["lab_a2"]
    orphan = PatientEncounters(
        name="Orphan", patient_id="ORPH-1", capture_date="2026-09-01", capture_date_dt=date(2026, 9, 1),
        lab_unit_id=lab_a2.id, project_id=None, disease_id=core_test_data["dr"].id, is_set_based=True,
    )
    db_session.add(orphan)
    db_session.flush()
    client = _client(app, db_session, world["admin"])
    base = "/api/admin/no-project/export.duckdb"
    assert client.get(base).status_code == 400
    assert client.get(base + "?lab_unit_id=abc").status_code == 400
    response = client.get(f"{base}?lab_unit_id={lab_a1.id}&lab_unit_id={lab_a2.id}&identifiers=1")
    assert response.status_code == 200
    con = _open(tmp_path, response.data)
    # only the project-less encounter: the project encounter in lab_a1 must not appear
    assert con.execute("select patient_id from encounters").fetchall() == [("ORPH-1",)]
    assert con.execute("select value from manifest where key = 'project_id'").fetchone() == ("",)
    # selected lab outside the selected hospital -> no labs -> 400
    mismatch = client.get(f"{base}?hospital_id={lab_a1.hospital_id + 9999}&lab_unit_id={lab_a1.id}")
    assert mismatch.status_code == 400


def test_admin_page_shell_fragment_and_options_api(app, db_session, world):
    assert _client(app, db_session, world["plain"]).get("/admin/project-export").status_code == 403
    assert _client(app, db_session, world["plain"]).get("/api/admin/project-export/options").status_code == 403
    admin = _client(app, db_session, world["admin"])
    shell = admin.get("/admin/project-export")
    assert shell.status_code == 200 and b"hx-get" in shell.data and b"/admin/project-export/workspace" in shell.data
    fragment = admin.get("/admin/project-export/workspace")
    assert fragment.status_code == 200
    assert b"ROUTE_P" in fragment.data and b"No project (classical uploads)" in fragment.data
    assert b'name="hospital_id" multiple' in fragment.data and b'name="lab_unit_id" multiple' in fragment.data
    options = admin.get("/api/admin/project-export/options").get_json()
    assert any(p["code"] == "ROUTE_P" for p in options["projects"])
    assert options["hospitals"] and options["lab_units"] and "hospital_id" in options["lab_units"][0]


def test_missing_engine_gives_503_not_a_crash(app, db_session, world, monkeypatch):
    import sys

    monkeypatch.setitem(sys.modules, "duckdb", None)
    response = _client(app, db_session, world["admin"]).get(f"/api/admin/projects/{world['project'].id}/export.duckdb")
    assert response.status_code == 503 and response.get_json()["error"] == "unavailable"


def test_excel_download_matches_the_package(app, db_session, world, tmp_path):
    from openpyxl import load_workbook

    client = _client(app, db_session, world["admin"])
    response = client.get(f"/api/admin/projects/{world['project'].id}/export.xlsx")
    assert response.status_code == 200
    assert response.headers["Content-Type"].startswith("application/vnd.openxmlformats")
    assert ".xlsx" in response.headers["Content-Disposition"]
    out = tmp_path / "x.xlsx"
    out.write_bytes(response.data)
    wb = load_workbook(out)
    assert "encounters" in wb.sheetnames and wb.sheetnames[-1] == "xlsx_notes"
    rows = list(wb["encounters"].iter_rows(values_only=True))
    assert rows[0][0] == "encounter_uuid" and rows[1][rows[0].index("patient_id")] is None
    assert client.get("/api/admin/no-project/export.xlsx").status_code == 400
    # switching user mid-test needs a fresh client (Flask-Login caches the user on g)
    assert _client(app, db_session, world["plain"]).get(
        f"/api/admin/projects/{world['project'].id}/export.xlsx").status_code == 403
