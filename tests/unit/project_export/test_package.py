from datetime import date

import duckdb
import pytest

from models import EncounterSetImage, PatientEncounters, Project
from project_export import ExportScope
from project_export.package import build_package
from project_export.schema import PACKAGE_TABLES, TABLES


@pytest.fixture
def world(db_session, core_test_data):
    lab = core_test_data["lab_a1"]
    project = Project(title="Export Project", code="EXP_P", active=True)
    db_session.add(project)
    db_session.flush()
    encounter = PatientEncounters(
        name="Jane Patient", patient_id="PID-1", capture_date="2026-09-14", capture_date_dt=date(2026, 9, 14),
        lab_unit_id=lab.id, project_id=project.id, disease_id=core_test_data["dr"].id, is_set_based=True,
        remarks="free text as entered",
        metadata_json={"patient": {"sex": "female", "patient_name": "Jane"}, "encounter": {"mode": "handheld"}},
    )
    db_session.add(encounter)
    db_session.flush()
    image = EncounterSetImage(
        patient_encounter_id=encounter.id, spatial_position=1, original_filename="img1.jpg",
        folder_rel="files/test", project_id=project.id, remarks="blurry", metadata_json={"laterality": "RIGHT"},
    )
    db_session.add(image)
    db_session.flush()
    return {"project": project, "lab": lab, "encounter": encounter, "image": image}


def _build(db_session, world, tmp_path, *, pii):
    scope = ExportScope(world["project"].id, frozenset({world["lab"].id}),
                        frozenset({world["lab"].id}) if pii else frozenset())
    path = tmp_path / ("pii.duckdb" if pii else "nopii.duckdb")
    counts = build_package(db_session, scope, path, project_code="EXP_P", project_title="Export Project")
    return path, counts


def test_package_without_identifiers_masks_patient_fields(db_session, world, tmp_path):
    path, counts = _build(db_session, world, tmp_path, pii=False)
    assert (counts["encounters"], counts["images"]) == (1, 1)
    con = duckdb.connect(str(path), read_only=True)
    enc = con.execute("select patient_id, patient_name, remarks, source_path, image_count, identifiers_included, "
                      "json_extract_string(metadata, '$.encounter.mode'), json_extract_string(metadata, '$.patient.patient_name') "
                      "from encounters").fetchone()
    assert enc == (None, None, "free text as entered", "other", 1, False, "handheld", None)
    img = con.execute("select kind, original_filename, remarks, json_extract_string(metadata, '$.laterality'), lab_unit_id "
                      "from images").fetchone()
    assert img == ("encounter_set_image", None, "blurry", "RIGHT", world["lab"].id)


def test_package_with_identifiers_and_dictionary(db_session, world, tmp_path):
    path, _ = _build(db_session, world, tmp_path, pii=True)
    con = duckdb.connect(str(path), read_only=True)
    assert con.execute("select patient_id, patient_name from encounters").fetchone() == ("PID-1", "Jane Patient")
    assert con.execute("select original_filename from images").fetchone() == ("img1.jpg",)
    manifest = dict(con.execute("select key, value from manifest").fetchall())
    assert manifest["schema_version"] == "1" and manifest["rows_encounters"] == "1"
    documented = {(t, c) for t, c in con.execute(
        "select table_name, column_name from data_dictionary where column_name <> '' and stage <> 'derived'").fetchall()}
    assert documented == {(t.name, c.name) for t in TABLES + PACKAGE_TABLES for c in t.columns}


def test_existing_file_is_not_overwritten(db_session, world, tmp_path):
    path, _ = _build(db_session, world, tmp_path, pii=False)
    scope = ExportScope(world["project"].id, frozenset({world["lab"].id}))
    with pytest.raises(FileExistsError):
        build_package(db_session, scope, path, project_code="X", project_title="X")


def test_metadata_is_also_flattened_into_typed_columns(db_session, world, tmp_path):
    world["encounter"].metadata_json = {"patient": {"age": 61, "monocular": True}, "encounter": {"mode": "handheld"}}
    db_session.flush()
    path, _ = _build(db_session, world, tmp_path, pii=False)
    con = duckdb.connect(str(path), read_only=True)
    cols = dict(con.execute("select column_name, type from data_dictionary where table_name = 'encounter_metadata' "
                            "and column_name <> ''").fetchall())
    assert cols["patient.age"] in {"UBIGINT", "BIGINT"} and cols["patient.monocular"] == "BOOLEAN"
    assert con.execute('select "patient.age", "encounter.mode" from encounter_metadata').fetchone() == (61, "handheld")
    assert con.execute("select value from manifest where key = 'typed_encounter_metadata'").fetchone() == ("built",)
    # images with metadata get the same treatment
    assert con.execute("select value from manifest where key = 'typed_image_metadata_values'").fetchone() == ("built",)


def test_awkward_metadata_never_breaks_the_export(db_session, world, tmp_path):
    # a top-level key equal to the join column collides in the flattened table (DuckDB renames or we skip)
    world["encounter"].metadata_json = {"encounter_uuid": "x"}
    db_session.flush()
    path, counts = _build(db_session, world, tmp_path, pii=False)
    assert counts["encounters"] == 1
    con = duckdb.connect(str(path), read_only=True)
    status = con.execute("select value from manifest where key = 'typed_encounter_metadata'").fetchone()[0]
    assert status == "built" or status.startswith("skipped:")
    assert con.execute("select count(*) from encounters").fetchone() == (1,)
    names = [row[0] for row in con.execute("describe encounter_metadata").fetchall()] if status == "built" else []
    assert len(names) == len(set(names))


def test_every_schema_table_exists_and_matches_its_columns(db_session, world, tmp_path):
    path, _ = _build(db_session, world, tmp_path, pii=False)
    con = duckdb.connect(str(path), read_only=True)
    for table in TABLES + PACKAGE_TABLES:
        cols = [row[0] for row in con.execute(f'describe "{table.name}"').fetchall()]
        assert cols == list(table.column_names), table.name


def test_guide_relationships_and_example_queries_are_valid(db_session, world, tmp_path):
    path, _ = _build(db_session, world, tmp_path, pii=False)
    con = duckdb.connect(str(path), read_only=True)
    columns = {t.name: set(t.column_names) for t in TABLES}
    dynamic = {"encounter_metadata", "image_metadata_values"}
    for from_table, from_col, to_table, to_col, _note in con.execute("select * from relationships").fetchall():
        for table, col in ((from_table, from_col), (to_table, to_col)):
            assert table in dynamic or col in columns[table], (table, col)
    queries = con.execute("select title, text from guide where section = 'example_query'").fetchall()
    assert queries
    for title, sql in queries:
        con.execute(sql).fetchall()  # must parse and run (empty results are fine)
    assert {row[0] for row in con.execute("select distinct section from guide").fetchall()} >= {
        "overview", "stages", "rules", "caveats", "example_query"}


def test_pii_flagged_image_keeps_its_rows_but_not_its_filename(db_session, world, tmp_path):
    from models import ImagePiiVerification

    db_session.add(ImagePiiVerification(image_uuid=world["image"].uuid, image_variant="orig", pii_status="detected"))
    db_session.flush()
    path, counts = _build(db_session, world, tmp_path, pii=False)
    con = duckdb.connect(str(path), read_only=True)
    assert con.execute("select pii_detected, original_filename from images").fetchone() == (True, None)
    assert con.execute("select count(*) from verification_images").fetchone() == (1,)


def _grade_world(db_session, world, core_test_data):
    from sqlalchemy import select

    from models import Consensus, DiseaseGrading, Grade, GradingTask
    from tests.helpers.factories import UserFactory

    label = db_session.execute(
        select(DiseaseGrading).where(DiseaseGrading.disease_id == core_test_data["dr"].id).limit(1)
    ).scalar_one()
    grader = UserFactory.create_admin(db_session, username="stage_grader")
    task = GradingTask(
        encounter_set_image_id=world["image"].id, disease_id=core_test_data["dr"].id, lab_unit_id=world["lab"].id,
        project_id=world["project"].id, grading_target_level="image", task_source="profile_package_negative_control",
        state="final",
    )
    db_session.add(task)
    db_session.flush()
    human = Grade(task_id=task.id, grader_user_id=grader.id, role_slot="resident", disease_grading_id=label.id,
                  grade_name=label.impression, comment="looks fine", time_taken=12.5)
    ai = Grade(task_id=task.id, grader_user_id=grader.id, role_slot="ai", disease_grading_id=label.id,
               grade_name=label.impression, ai_model_name="wai_test", ai_model_version="v1",
               ai_review_status="minor_miss", ai_review_comment="slightly off")
    db_session.add_all([human, ai])
    db_session.add(Consensus(task_id=task.id, final_disease_grading_id=label.id, method="match",
                             final_grade_name=label.impression, final_disease_name="DR"))
    db_session.flush()
    return task, label


def test_tasks_grades_and_consensus_stage(db_session, world, core_test_data, tmp_path):
    task, label = _grade_world(db_session, world, core_test_data)
    world["image"].is_not_gradable = True
    world["image"].not_gradable_reason = "too dark"
    db_session.flush()
    path, counts = _build(db_session, world, tmp_path, pii=False)
    assert (counts["tasks"], counts["grades"], counts["consensus"]) == (1, 2, 1)
    con = duckdb.connect(str(path), read_only=True)
    assert con.execute("select target_level, task_source, is_negative_control, state, grade_count, "
                       "image_uuid = ?, encounter_uuid = ? from tasks", [world["image"].uuid, world["encounter"].uuid]
                       ).fetchone() == ("image", "profile_package_negative_control", True, "final", 2, True, True)
    rows = dict(con.execute("select role_slot, grade_name from grades").fetchall())
    assert rows == {"resident": label.impression, "ai": label.impression}
    ai = con.execute("select ai_model, ai_model_version, ai_review_status, ai_review_comment from grades "
                     "where role_slot = 'ai'").fetchone()
    assert ai == ("wai_test", "v1", "minor_miss", "slightly off")
    assert con.execute("select comment, time_taken_seconds from grades where role_slot = 'resident'").fetchone() == (
        "looks fine", 12.5)
    assert con.execute("select method, final_grade from consensus").fetchone() == ("match", label.impression)
    # the same grade ids resolve through the lookup tables and verification stage
    assert con.execute("select count(*) from dim_grade_labels where grade_name = ?", [label.impression]).fetchone()[0] >= 1
    assert con.execute("select is_not_gradable, not_gradable_reason from verification_images").fetchone() == (True, "too dark")


def test_side_normalisation():
    from project_export.stage_verification import side

    assert [side(v) for v in ("RIGHT", "od", "R", "Left", "OS", "ou", None, "", "weird")] == [
        "R", "R", "R", "L", "L", "B", None, None, None]


def test_coco_rows_keep_valid_boxes_and_flag_the_rest(db_session, world, core_test_data, tmp_path):
    from sqlalchemy import select

    from grading.workbench.models import AnnotationInstance, AnnotationSet
    from models import Grade, GradingsFeatures

    task, label = _grade_world(db_session, world, core_test_data)
    grade = db_session.execute(select(Grade).where(Grade.task_id == task.id, Grade.role_slot == "resident")).scalar_one()
    feature = GradingsFeatures(disease_grading_id=label.id, sr_no=9101, label="Export lesion")
    db_session.add(feature)
    db_session.flush()
    annotation_set = AnnotationSet(grade_id=grade.id, policy_source="project", policy_revision=1,
                                   source_image_width=100, source_image_height=80)
    db_session.add(annotation_set)
    db_session.flush()

    def instance(order, **bbox):
        db_session.add(AnnotationInstance(
            annotation_set_id=annotation_set.id, image_uuid=world["image"].uuid, class_source="grading_feature",
            grading_feature_id=feature.id, class_key_snapshot="lesion", class_label_snapshot="Lesion",
            policy_revision=1, geometry_type="box", geometry_json={}, instance_order=order, **bbox,
        ))

    instance(0, bbox_x=10, bbox_y=10, bbox_w=20, bbox_h=15)       # valid
    instance(1, bbox_x=95, bbox_y=70, bbox_w=20, bbox_h=15)       # runs past the image edge -> invalid
    db_session.flush()
    path, counts = _build(db_session, world, tmp_path, pii=False)
    assert counts["coco_annotations"] == 2 and counts["coco_categories"] == 1
    con = duckdb.connect(str(path), read_only=True)
    rows = con.execute("select bbox_x, bbox_y, bbox_w, bbox_h, area, width, height, category_name, source_role_slot, "
                       "source_grade_id, warning, category_id from coco_annotations order by annotation_id").fetchall()
    assert rows[0] == (10.0, 10.0, 20.0, 15.0, 300.0, 100, 80, "lesion", "resident", grade.id, None, 1)
    assert rows[1][0] is None and rows[1][10] == "invalid annotation geometry"
    assert con.execute("select count(*), sum(annotation_count) from coco_images").fetchone() == (1, 2)
    assert con.execute("select category_name, annotation_count from coco_categories").fetchone() == ("lesion", 2)


def test_timing_and_review_tables_are_present_and_empty_without_data(db_session, world, core_test_data, tmp_path):
    _grade_world(db_session, world, core_test_data)
    path, counts = _build(db_session, world, tmp_path, pii=False)
    for name in ("grading_sessions", "grading_session_targets", "review_history", "regrade_tasks", "review_corrections"):
        assert counts[name] == 0


def test_workbook_is_generated_from_the_duckdb_file_and_reports_limits(db_session, world, tmp_path):
    from openpyxl import load_workbook

    from project_export.xlsx import write_workbook

    world["encounter"].remarks = "x" * 50
    db_session.flush()
    path, _ = _build(db_session, world, tmp_path, pii=False)
    out = tmp_path / "out.xlsx"
    notes = write_workbook(path, out, max_rows=1, cell_limit=20)
    wb = load_workbook(out)  # full mode: a write-only file has no stored dimensions, which makes read-only mode crawl
    names = wb.sheetnames
    assert names[:4] == ["manifest", "guide", "relationships", "data_dictionary"] and names[-1] == "xlsx_notes"
    assert {t.name for t in TABLES} <= set(names)
    rows = list(wb["encounters"].iter_rows(values_only=True))
    header, first = rows[0], rows[1]
    assert first[header.index("remarks")] == "x" * 20            # truncated at the cell limit
    assert first[header.index("capture_date")] is not None       # dates survive
    assert any("truncated" in note for _, note in notes) and any("first 1 of" in note for _, note in notes)
    assert list(wb["xlsx_notes"].iter_rows(values_only=True))[0] == ("table", "note")
    assert len(list(wb["guide"].iter_rows())) == 2               # header + first row only (max_rows=1)


def test_workbook_handles_list_valued_metadata(db_session, world, tmp_path):
    from openpyxl import load_workbook

    from project_export.xlsx import write_workbook

    world["encounter"].metadata_json = {"encounter": {"captured_positions": ["od", "os"], "mode": "handheld"}}
    db_session.flush()
    path, _ = _build(db_session, world, tmp_path, pii=False)
    out = tmp_path / "lists.xlsx"
    write_workbook(path, out)
    sheet = load_workbook(out)["encounter_metadata"]
    header, first = [c.value for c in sheet[1]], [c.value for c in sheet[2]]
    assert first[header.index("encounter.captured_positions")] == '["od","os"]'


def test_ai_payloads_drop_identifying_keys_without_a_grant(db_session, world, core_test_data, tmp_path):
    from models import AIInferenceRun, AIModel

    task, _ = _grade_world(db_session, world, core_test_data)
    model = AIModel(name="m", version="1")
    db_session.add(model)
    db_session.flush()
    db_session.add(AIInferenceRun(
        task_id=task.id, ai_model_id=model.id, source="internal", status="success",
        request_manifest_json={"filename": "jane_doe.jpg", "sex": "female"},
        execute_response_json={"results": [{"filename": "jane_doe.jpg", "prediction": "refer"}]},
    ))
    db_session.flush()
    path, _ = _build(db_session, world, tmp_path, pii=False)
    con = duckdb.connect(str(path), read_only=True)
    request, response = con.execute("select request_manifest, response from ai_runs").fetchone()
    assert "jane_doe" not in request and "jane_doe" not in response and "refer" in response and "female" in request
    path2, _ = _build(db_session, world, tmp_path, pii=True)
    con2 = duckdb.connect(str(path2), read_only=True)
    assert "jane_doe.jpg" in con2.execute("select response from ai_runs").fetchone()[0]


def test_file_names_in_error_text_are_redacted_without_a_grant():
    from project_export.tables import redact_files

    text = "Failed to read 'jane_doe 12.JPG' and /tmp/scan.final.pdf, then raised ValueError"
    assert redact_files(text, pii=False) == "Failed to read '[file]' and /tmp/[file], then raised ValueError"
    assert redact_files(text, pii=True) == text and redact_files(None, pii=False) is None


def test_intra_rater_tables_follow_the_source_task(db_session, world, core_test_data, tmp_path):
    from models import IntraRaterBatch, IntraRaterGrade, IntraRaterTask
    from tests.helpers.factories import UserFactory

    task, label = _grade_world(db_session, world, core_test_data)
    grader = UserFactory.create_admin(db_session, username="intra_grader")
    batch = IntraRaterBatch(disease_id=core_test_data["dr"].id, lab_unit_id=world["lab"].id, target_images_per_grader=5,
                            normal_grade_id=label.id, selection_snapshot_json='{"strategy": "random"}', remarks="pilot")
    db_session.add(batch)
    db_session.flush()
    repeat = IntraRaterTask(batch_id=batch.id, grader_user_id=grader.id, disease_id=core_test_data["dr"].id,
                            lab_unit_id=world["lab"].id, source_task_id=task.id, state="graded")
    db_session.add(repeat)
    db_session.flush()
    db_session.add(IntraRaterGrade(task_id=repeat.id, batch_id=batch.id, grader_user_id=grader.id,
                                   disease_grading_id=label.id, grade_name=label.impression, comment="again",
                                   selected_features_json='[1, 2]'))
    db_session.flush()
    path, counts = _build(db_session, world, tmp_path, pii=False)
    assert (counts["intra_rater_batches"], counts["intra_rater_tasks"], counts["intra_rater_grades"]) == (1, 1, 1)
    con = duckdb.connect(str(path), read_only=True)
    assert con.execute("select source_task_uuid = ?, state, grader_username from intra_rater_tasks", [task.uuid]).fetchone() == (
        True, "graded", "intra_grader")
    assert con.execute("select normal_grade, remarks, json_extract_string(selection_snapshot, '$.strategy') "
                       "from intra_rater_batches").fetchone() == (label.impression, "pilot", "random")
    assert con.execute("select grade_name, comment, selected_features from intra_rater_grades").fetchone() == (
        label.impression, "again", "[1, 2]")
